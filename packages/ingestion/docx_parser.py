"""DOCX parser extracting paragraphs, numbered clauses, and tables."""

from __future__ import annotations

import io
from pathlib import Path

import docx

from packages.extraction.regex_extractor import detect_language
from packages.ingestion.schema import ClauseSegment, ParsedDocument


class DocxParser:
    """Parses DOCX documents into structured clauses with paragraph and table indexing."""

    def parse_file(self, file_path: Path) -> ParsedDocument:
        with open(file_path, "rb") as f:
            return self.parse_bytes(f.read(), filename=file_path.name)

    def parse_bytes(self, content: bytes, filename: str = "document.docx") -> ParsedDocument:
        doc = docx.Document(io.BytesIO(content))

        clauses: list[ClauseSegment] = []
        full_text_parts: list[str] = []
        curr_offset = 0
        clause_counter = 1

        # 1. Extract regular paragraphs
        for p_idx, para in enumerate(doc.paragraphs, start=1):
            text = para.text.strip()
            if not text:
                continue

            start_char = curr_offset
            end_char = start_char + len(text)
            curr_offset = end_char + 2  # account for double newline

            full_text_parts.append(text)
            lang = detect_language(text)

            clauses.append(
                ClauseSegment(
                    id=f"docx_p{p_idx}",
                    text=text,
                    paragraph_index=p_idx,
                    start_char=start_char,
                    end_char=end_char,
                    language=lang,
                    source_type="docx_paragraph",
                )
            )
            clause_counter += 1

        # 2. Extract tables (each row as a potential specification clause)
        for t_idx, table in enumerate(doc.tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if not row_cells:
                    continue

                row_text = " | ".join(row_cells)
                start_char = curr_offset
                end_char = start_char + len(row_text)
                curr_offset = end_char + 2

                full_text_parts.append(row_text)
                lang = detect_language(row_text)

                clauses.append(
                    ClauseSegment(
                        id=f"docx_t{t_idx}_r{r_idx}",
                        text=row_text,
                        paragraph_index=clause_counter,
                        start_char=start_char,
                        end_char=end_char,
                        language=lang,
                        source_type="docx_table",
                    )
                )
                clause_counter += 1

        full_doc_text = "\n\n".join(full_text_parts)

        return ParsedDocument(
            filename=filename,
            source_type="docx",
            full_text=full_doc_text,
            clauses=clauses,
            page_count=1,
            is_ocr=False,
            status="SUCCESS",
        )
