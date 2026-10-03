from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

import docx
from docx.shared import Pt, RGBColor

BANNER_TEXT = "This tool flags issues for the officer to review. It does not approve or reject a tender."


def generate_docx_report(
    audit_meta: dict[str, Any],
    clauses: list[dict[str, Any]],
    findings: list[dict[str, Any]],
) -> bytes:
    """Generate official bilingual tender audit report in DOCX format."""
    doc = docx.Document()

    # Document Title
    title = doc.add_heading("Tender Linter: Specification Audit Report", level=0)
    title.runs[0].font.size = Pt(20)
    title.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    # Mandatory Reviewer Banner
    banner_p = doc.add_paragraph()
    banner_run = banner_p.add_run(f"NOTICE: {BANNER_TEXT}")
    banner_run.bold = True
    banner_run.font.size = Pt(10)
    banner_run.font.color.rgb = RGBColor(180, 83, 9)

    doc.add_paragraph()

    # Metadata Grid Table
    doc_name = audit_meta.get("document_name") or "Unnamed Document"
    audit_id = str(audit_meta.get("id", ""))
    created_at = audit_meta.get("created_at") or datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    lang_hint = str(audit_meta.get("language_hint", "en")).upper()

    counts = {"ERROR": 0, "WARNING": 0, "INFO": 0, "CANNOT_VERIFY": 0}
    for f in findings:
        sev = f.get("severity", "INFO")
        counts[sev] = counts.get(sev, 0) + 1

    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.style = "Table Grid"

    rows = meta_table.rows
    rows[0].cells[0].paragraphs[0].add_run("Document Name: ").bold = True
    rows[0].cells[0].paragraphs[0].add_run(doc_name)
    rows[0].cells[1].paragraphs[0].add_run("Audit ID: ").bold = True
    rows[0].cells[1].paragraphs[0].add_run(audit_id)

    rows[1].cells[0].paragraphs[0].add_run("Audited On: ").bold = True
    rows[1].cells[0].paragraphs[0].add_run(str(created_at))
    rows[1].cells[1].paragraphs[0].add_run("Language Hint: ").bold = True
    rows[1].cells[1].paragraphs[0].add_run(lang_hint)

    rows[2].cells[0].paragraphs[0].add_run("Total Clauses: ").bold = True
    rows[2].cells[0].paragraphs[0].add_run(str(len(clauses)))
    rows[2].cells[1].paragraphs[0].add_run("Findings: ").bold = True
    rows[2].cells[1].paragraphs[0].add_run(
        f"{len(findings)} (Errors: {counts['ERROR']}, Warnings: {counts['WARNING']}, Info: {counts['INFO']})"
    )

    doc.add_paragraph()
    doc.add_heading("Audit Findings and Verified Evidence", level=1)

    if not findings:
        doc.add_paragraph("No standard discrepancies or defect rules triggered.")
    else:
        clause_map = {c.get("id"): c.get("text", "") for c in clauses}

        findings_table = doc.add_table(rows=1, cols=5)
        findings_table.style = "Table Grid"

        hdr_cells = findings_table.rows[0].cells
        headers = ["Clause / Rule", "Severity", "Explanation (EN / HI)", "Evidence & Verification", "Officer Action"]
        for i, text in enumerate(headers):
            p = hdr_cells[i].paragraphs[0]
            r = p.add_run(text)
            r.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)

        for f in findings:
            row_cells = findings_table.add_row().cells
            cid = f.get("clause_id", "")
            clause_snip = clause_map.get(cid, "")
            if len(clause_snip) > 80:
                clause_snip = clause_snip[:77] + "..."

            rule_id = f.get("rule_id", "")
            sev = f.get("severity", "INFO")
            msg_en = f.get("message_en", "")
            msg_hi = f.get("message_hi", "")

            ev = f.get("evidence") or {}
            ev_url = ev.get("url") or "N/A"
            ev_date = ev.get("verified_on") or "N/A"
            ev_ref = ev.get("evidence_ref") or ""

            # Col 0: Clause and Rule
            p0 = row_cells[0].paragraphs[0]
            p0.add_run(f"Clause: {cid}\n").bold = True
            p0.add_run(f'"{clause_snip}"\n')
            p0.add_run(f"Rule: {rule_id}").bold = True

            # Col 1: Severity
            p1 = row_cells[1].paragraphs[0]
            p1.add_run(sev).bold = True

            # Col 2: Explanation
            p2 = row_cells[2].paragraphs[0]
            p2.add_run("[EN] ").bold = True
            p2.add_run(f"{msg_en}\n\n")
            p2.add_run("[HI] ").bold = True
            p2.add_run(msg_hi)

            # Col 3: Evidence
            p3 = row_cells[3].paragraphs[0]
            p3.add_run(f"Verified Date: {ev_date}\n").bold = True
            p3.add_run("Evidence URL: ").bold = True
            p3.add_run(f"{ev_url}\n")
            if ev_ref:
                p3.add_run(f"Ref: {ev_ref}")

            # Col 4: Action
            p4 = row_cells[4].paragraphs[0]
            decision = f.get("decision") or "Pending Review"
            reason = f.get("decision_reason") or ""
            p4.add_run(f"Status: {decision}\n").bold = True
            if reason:
                p4.add_run(f"Reason: {reason}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
